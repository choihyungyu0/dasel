"""docs/다셀_기능명세서.xlsx 의 모든 시트를 docs/spec/<시트명>.md 로 변환한다."""
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "다셀_기능명세서.xlsx"
OUT = ROOT / "docs" / "spec"


def cell(v):
    if v is None:
        return ""
    return str(v).strip().replace("|", r"\|").replace("\r", "").replace("\n", "<br>")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC, data_only=True)
    for ws in wb:
        rows = [[cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        used = [i for i in range(ws.max_column) if any(i < len(r) and r[i] for r in rows)]
        rows = [[r[i] for i in used] for r in rows]
        lines = [f"# {ws.title}", "", f"> 원본: docs/{SRC.name} · 시트 「{ws.title}」 · {len(rows)}행", ""]
        head, body = rows[0], rows[1:]
        lines.append("| " + " | ".join(head) + " |")
        lines.append("|" + "---|" * len(head))
        lines += ["| " + " | ".join(r) + " |" for r in body]
        (OUT / f"{ws.title}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
        print(f"{ws.title}.md  rows={len(rows)} cols={len(used)}")


if __name__ == "__main__":
    main()
